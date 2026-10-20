"""QC Form Generator for Viedoc EDC.

Populates GCP/SOP-compliant Excel QC forms from parsed Viedoc Configuration Reports,
handling page headers, cell formatting, column widths, autofilters, Critical Attributes,
and 6-column QC tracking schema.
"""

from copy import copy
from datetime import datetime
from pathlib import Path
import re
from typing import Any, List, Optional, Tuple
import zipfile
import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from src.critical_attributes import evaluate_row_critical_attributes
from src.models import (
    ParsedSheet,
    ParsedViedocConfig,
    SHEET_CONFIGS,
    SHEET_CONFIGS_BY_QC_NAME,
    SheetConfig,
)
from src.parser import (
    format_header_date,
    generate_output_filename,
    update_header_footer_placeholders,
)

SHEET_HEADER_TITLES = {
    "Viedoc - Data Checks": "Data Checks",
    "Viedoc - Alerts": "Alerts",
    "Viedoc - Items and Groups": "Items and Groups",
    "Viedoc - Events": "Study Workflow - Events",
    "Viedoc - Activities": "Study Workflow - Activities",
    "Viedoc - Forms": "Study Workflow - Forms",
    "Viedoc - Functions": "Functions and Conditions",
    "Viedoc - Code lists": "Code lists",
    "Amendment Change Log": "Amendment Change Log",
}

DEFAULT_VALIDATION_TEMPLATE = "GCPS_DMF_012.8_A-CValidation Check QC Form_DRAFT.xlsx"
DEFAULT_DESIGN_TEMPLATE = "GCPS_DMF_026.5_A-EDesign and Specification QC Form_DRAFT.xlsm"

VALIDATION_TEMPLATE_KEYWORD = "validation check qc form"
DESIGN_TEMPLATE_KEYWORD = "design and specification qc form"

# The 6 standard trailing QC columns per SOI-EDC-QC-001 Section 3.2
QC_TRAILING_COLUMNS = [
    "Critical Attributes",
    "QC Method",
    "Validation Date",
    "Validator",
    "Validation Status",
    "Comment",
]


def find_template_path(templates_dir: Path, qc_form_type: str) -> Path:
    """Dynamically locates a template file in templates_dir based on QC form type."""
    if not templates_dir.exists():
        raise FileNotFoundError(f"Templates directory does not exist: {templates_dir}")

    form_type = qc_form_type.lower().strip()
    if form_type == "validation":
        keyword = VALIDATION_TEMPLATE_KEYWORD
        default_name = DEFAULT_VALIDATION_TEMPLATE
    elif form_type == "design":
        keyword = DESIGN_TEMPLATE_KEYWORD
        default_name = DEFAULT_DESIGN_TEMPLATE
    else:
        raise ValueError(
            f"Unknown qc_form_type: '{qc_form_type}'. Expected 'validation' or 'design'."
        )

    candidates = []
    for f in templates_dir.iterdir():
        if f.is_file() and not f.name.startswith("~$"):
            suffix = f.suffix.lower()
            if suffix in [".xlsx", ".xlsm"] and keyword in f.name.lower():
                candidates.append(f)

    if candidates:
        candidates.sort(key=lambda p: p.stat().st_mtime, reverse=True)
        return candidates[0]

    fallback = templates_dir / default_name
    if fallback.exists():
        return fallback

    raise FileNotFoundError(
        f"No template found matching keyword '{keyword}' (*.xlsx/*.xlsm) in {templates_dir}"
    )


DATA_FONT = Font(name="Arial", size=10, bold=False, color="000000")
BORDER_SIDE = Side(style="thin", color="D9D9D9")
DATA_BORDER = Border(
    left=BORDER_SIDE, right=BORDER_SIDE, top=BORDER_SIDE, bottom=BORDER_SIDE
)
DATA_ALIGNMENT = Alignment(vertical="top", wrap_text=True)

QC_CRITICAL_ATTR_COLUMN_WIDTH = 28
QC_METHOD_COLUMN_WIDTH = 16
QC_STATUS_COLUMN_WIDTH = 18
QC_DATE_COLUMN_WIDTH = 15
QC_VALIDATOR_COLUMN_WIDTH = 16
QC_COMMENT_COLUMN_WIDTH = 28


def ensure_six_qc_headers(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    spec_col_count: int,
    has_design_version_col: bool,
) -> int:
    """Ensures Row 1 has the 6 standard trailing QC headers after the specification columns,
    copying header styling from the existing template QC header cells.
    Returns total column count.
    """
    qc_start_col = (1 if has_design_version_col else 0) + spec_col_count + 1
    style_ref_cell = ws.cell(1, max(1, qc_start_col))

    for i, col_title in enumerate(QC_TRAILING_COLUMNS):
        col_idx = qc_start_col + i
        cell = ws.cell(1, col_idx)
        cell.value = col_title
        if style_ref_cell.has_style:
            cell.font = copy(style_ref_cell.font)
            cell.fill = copy(style_ref_cell.fill)
            cell.border = copy(style_ref_cell.border)
            cell.alignment = copy(style_ref_cell.alignment)

    total_cols = qc_start_col + len(QC_TRAILING_COLUMNS) - 1
    if ws.max_column > total_cols:
        ws.delete_cols(total_cols + 1, ws.max_column - total_cols)

    return total_cols


def adjust_column_widths(
    ws: openpyxl.worksheet.worksheet.Worksheet, total_cols: int, total_rows: int
) -> None:
    """Intelligently adjusts column widths based on header and data content length."""
    sample_row_limit = min(total_rows, 50)

    for col_idx in range(1, total_cols + 1):
        col_letter = get_column_letter(col_idx)
        header_val = ws.cell(1, col_idx).value
        header_text = str(header_val) if header_val is not None else ""

        if "Critical Attributes" in header_text:
            ws.column_dimensions[col_letter].width = QC_CRITICAL_ATTR_COLUMN_WIDTH
            continue
        elif "QC Method" in header_text or "Subject ID" in header_text:
            ws.column_dimensions[col_letter].width = QC_METHOD_COLUMN_WIDTH
            continue
        elif "Validation Status" in header_text:
            ws.column_dimensions[col_letter].width = QC_STATUS_COLUMN_WIDTH
            continue
        elif "Validation Date" in header_text:
            ws.column_dimensions[col_letter].width = QC_DATE_COLUMN_WIDTH
            continue
        elif "Validator" in header_text:
            ws.column_dimensions[col_letter].width = QC_VALIDATOR_COLUMN_WIDTH
            continue
        elif "Comment" in header_text:
            ws.column_dimensions[col_letter].width = QC_COMMENT_COLUMN_WIDTH
            continue

        max_len = len(header_text)
        for row_idx in range(2, sample_row_limit + 1):
            val = ws.cell(row_idx, col_idx).value
            if val is not None:
                first_line = str(val).split("\n")[0]
                max_len = max(max_len, len(first_line))

        calculated_width = min(max(max_len + 3, 12), 45)
        ws.column_dimensions[col_letter].width = calculated_width


def update_sheet_header_footer(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    study_code: str,
    design_version: str,
    header_date: str,
    qc_form_name: Optional[str] = None,
) -> None:
    """Updates page header and footer in the worksheet adhering to Issue #5 / templates:
    - Middle cell: 4 separate lines with Arial sizes 16, 14, 11, 11 (middle aligned):
        Line 1: QC Form Name
        Line 2: Sheet Name
        Line 3: Version No.: <X.X>
        Line 4: Version Date: <DD/MMM/YYYY>
    - Left cell: &G (GCP-Service logo)
    - Right cell: Study Code & System
    """
    if not qc_form_name:
        wb = ws.parent
        sheetnames = wb.sheetnames if wb else [ws.title]
        if any("Data Checks" in s or "Alerts" in s for s in sheetnames):
            qc_form_name = "Validation Check QC Form"
        else:
            qc_form_name = "Design and Specification QC Form"

    sheet_title = SHEET_HEADER_TITLES.get(ws.title, ws.title.replace("Viedoc - ", ""))

    center_text = (
        f'&"Arial"&16{qc_form_name}\n'
        f'&"Arial"&14{sheet_title}\n'
        f'&"Arial"&11Version No.: {design_version}\n'
        f'&"Arial"&11Version Date: {header_date}'
    )
    left_text = "&G"
    right_text = (
        f'&"Arial"&10Study Code: {study_code}\n'
        f"System: Viedoc"
    )

    hf = ws.HeaderFooter
    for section in [hf.oddHeader, hf.evenHeader, hf.firstHeader]:
        if section is not None:
            section.color = None
            section.left.color = None
            section.left.text = left_text
            section.center.color = None
            section.center.text = center_text
            section.right.color = None
            section.right.text = right_text

    for section in [hf.oddFooter, hf.evenFooter, hf.firstFooter]:
        if section is not None:
            section.color = None
            for part in [section.left, section.center, section.right]:
                if part:
                    part.color = None
                    if part.text:
                        part.text = update_header_footer_placeholders(
                            part.text, study_code, design_version, header_date
                        )


def inject_header_logo_and_formatting(xlsx_path: Path, template_path: Path) -> None:
    """Post-processes an openpyxl-saved xlsx file to restore the GCP-Service logo in the
    Left Header via VML legacyDrawingHF, ensures Content_Types and sheet relationships
    are completely valid, and normalizes line breaks in headers.
    """
    logo_bytes = None
    logo_filename = "image1.jpeg"
    logo_ext = ".jpeg"

    search_paths = [template_path]
    templates_dir = template_path.parent
    if templates_dir.exists():
        for f in templates_dir.glob("*.xlsx"):
            if f not in search_paths:
                search_paths.append(f)

    for p in search_paths:
        if p.exists():
            try:
                with zipfile.ZipFile(p, "r") as tz:
                    for name in tz.namelist():
                        if "image1.jpeg" in name or (
                            "media" in name and name.endswith((".jpeg", ".jpg", ".png"))
                        ):
                            logo_bytes = tz.read(name)
                            logo_ext = Path(name).suffix.lower()
                            logo_filename = f"image1{logo_ext}"
                            break
            except Exception:
                pass
        if logo_bytes:
            break

    if not logo_bytes:
        return

    entries = {}
    with zipfile.ZipFile(xlsx_path, "r") as z:
        for item in z.infolist():
            entries[item.filename] = z.read(item.filename)

    # 1. Ensure xl/media/<logo_filename> exists
    media_path = f"xl/media/{logo_filename}"
    entries[media_path] = logo_bytes

    # 2. Ensure Content_Types has vml and image extension defaults
    ct_xml = entries["[Content_Types].xml"].decode("utf-8")
    modified_ct = False
    if 'Extension="vml"' not in ct_xml:
        vml_def = '<Default Extension="vml" ContentType="application/vnd.openxmlformats-officedocument.vmlDrawing"/>'
        ct_xml = ct_xml.replace("</Types>", f"{vml_def}</Types>")
        modified_ct = True
    img_ext = logo_ext.lstrip(".")
    img_content_type = "image/jpeg" if img_ext in ("jpeg", "jpg") else f"image/{img_ext}"
    if f'Extension="{img_ext}"' not in ct_xml:
        img_def = f'<Default Extension="{img_ext}" ContentType="{img_content_type}"/>'
        ct_xml = ct_xml.replace("</Types>", f"{img_def}</Types>")
        modified_ct = True
    if modified_ct:
        entries["[Content_Types].xml"] = ct_xml.encode("utf-8")

    # 3. Find all sheet XML files
    sheet_files = [f for f in entries.keys() if re.match(r"^xl/worksheets/sheet\d+\.xml$", f)]

    vml_template = """<xml xmlns:v="urn:schemas-microsoft-com:vml"
 xmlns:o="urn:schemas-microsoft-com:office:office"
 xmlns:x="urn:schemas-microsoft-com:excel">
 <o:shapelayout v:ext="edit">
  <o:idmap v:ext="edit" data="{data_id}"/>
 </o:shapelayout><v:shapetype id="_x0000_t75" coordsize="21600,21600" o:spt="75"
  o:preferrelative="t" path="m@4@5l@4@11@9@11@9@5xe" filled="f" stroked="f">
  <v:stroke joinstyle="miter"/>
  <v:formulas>
   <v:f eqn="if lineDrawn pixelLineWidth 0"/>
   <v:f eqn="sum @0 1 0"/>
   <v:f eqn="sum 0 0 @1"/>
   <v:f eqn="prod @2 1 2"/>
   <v:f eqn="prod @3 21600 pixelWidth"/>
   <v:f eqn="prod @3 21600 pixelHeight"/>
   <v:f eqn="sum @0 0 1"/>
   <v:f eqn="prod @6 1 2"/>
   <v:f eqn="prod @7 21600 pixelWidth"/>
   <v:f eqn="sum @8 21600 0"/>
   <v:f eqn="prod @7 21600 pixelHeight"/>
   <v:f eqn="sum @10 21600 0"/>
  </v:formulas>
  <v:path o:extrusionok="f" gradientshapeok="t" o:connecttype="rect"/>
  <o:lock v:ext="edit" aspectratio="t"/>
 </v:shapetype><v:shape id="LH" o:spid="_x0000_s{spid}" type="#_x0000_t75"
  style='position:absolute;margin-left:0;margin-top:0;width:167.4pt;height:39.6pt;
  z-index:1'>
  <v:imagedata o:relid="rId1" o:title="gcp_logo_60mm"/>
  <o:lock v:ext="edit" rotation="t"/>
 </v:shape></xml>"""

    vml_rels_content = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/{logo_filename}"/>
</Relationships>""".encode("utf-8")

    for sheet_path in sheet_files:
        sheet_xml = entries[sheet_path].decode("utf-8")
        sheet_num = re.search(r"sheet(\d+)\.xml", sheet_path).group(1)
        rels_path = f"xl/worksheets/_rels/sheet{sheet_num}.xml.rels"

        if rels_path in entries:
            rels_xml = entries[rels_path].decode("utf-8")
        else:
            rels_xml = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>'

        existing_rids = [int(m) for m in re.findall(r'Id="rId(\d+)"', rels_xml)]
        next_rid_num = max(existing_rids) + 1 if existing_rids else 1
        hf_rid = f"rIdHF{next_rid_num}"

        vml_filename = f"vmlDrawingHF{sheet_num}.vml"
        vml_path = f"xl/drawings/{vml_filename}"
        vml_rels_path = f"xl/drawings/_rels/{vml_filename}.rels"

        hf_rel_entry = f'<Relationship Id="{hf_rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/vmlDrawing" Target="../drawings/{vml_filename}"/>'
        rels_xml = rels_xml.replace("</Relationships>", f"{hf_rel_entry}</Relationships>")
        entries[rels_path] = rels_xml.encode("utf-8")

        data_id = 40 + int(sheet_num)
        spid = data_id * 1024 + 1
        entries[vml_path] = vml_template.format(data_id=data_id, spid=spid).encode("utf-8")
        entries[vml_rels_path] = vml_rels_content

        def clean_newlines(m):
            return m.group(0).replace("_x000a_", "\n")

        sheet_xml = re.sub(
            r"<(oddHeader|evenHeader|firstHeader)>.*?</\1>",
            clean_newlines,
            sheet_xml,
            flags=re.DOTALL,
        )

        if "<legacyDrawingHF" not in sheet_xml:
            hf_tag = f'<legacyDrawingHF xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{hf_rid}"/>'
            sheet_xml = sheet_xml.replace("</worksheet>", f"{hf_tag}</worksheet>")
            entries[sheet_path] = sheet_xml.encode("utf-8")

    temp_zip = xlsx_path.with_suffix(".tmp_zip")
    with zipfile.ZipFile(temp_zip, "w", compression=zipfile.ZIP_DEFLATED) as oz:
        for fpath, data in entries.items():
            oz.writestr(fpath, data)

    temp_zip.replace(xlsx_path)


def set_cell_value_safe(cell: openpyxl.cell.cell.Cell, value: Any) -> None:
    """Assigns value to a cell, ensuring strings starting with '=' are stored as plain text ('s')."""
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def populate_worksheet(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    parsed_sheet: ParsedSheet,
    cfg: SheetConfig,
    full_version: str,
) -> None:
    """Fills a single worksheet with parsed rows, Critical Attributes, and QC Method."""
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    total_cols = ensure_six_qc_headers(
        ws, len(parsed_sheet.lowest_headers), cfg.has_design_version_col
    )
    current_row = 2

    for row_vals in parsed_sheet.rows:
        col_offset = 1

        if cfg.has_design_version_col:
            cell = ws.cell(current_row, col_offset)
            set_cell_value_safe(cell, full_version)
            cell.font = DATA_FONT
            cell.border = DATA_BORDER
            cell.alignment = DATA_ALIGNMENT
            col_offset += 1

        for val in row_vals:
            cell = ws.cell(current_row, col_offset)
            set_cell_value_safe(cell, val)
            cell.font = DATA_FONT
            cell.border = DATA_BORDER
            cell.alignment = DATA_ALIGNMENT
            col_offset += 1

        # Evaluate Critical Attributes & default QC Method (CR-02, CR-03)
        crit_attrs, default_qc_method = evaluate_row_critical_attributes(
            cfg.qc_sheet_name, parsed_sheet.lowest_headers, row_vals
        )

        trailing_vals = [
            crit_attrs,         # Critical Attributes
            default_qc_method,  # QC Method ('Code review' or None)
            None,               # Validation Date
            None,               # Validator
            None,               # Validation Status
            None,               # Comment
        ]
        for t_val in trailing_vals:
            cell = ws.cell(current_row, col_offset)
            set_cell_value_safe(cell, t_val)
            cell.font = DATA_FONT
            cell.border = DATA_BORDER
            cell.alignment = DATA_ALIGNMENT
            col_offset += 1

        current_row += 1

    last_row = max(current_row - 1, 1)
    last_col_letter = get_column_letter(total_cols)
    ws.auto_filter.ref = f"A1:{last_col_letter}{last_row}"
    adjust_column_widths(ws, total_cols, last_row)


def populate_qc_form(
    parsed_config: ParsedViedocConfig,
    template_path: Path,
    output_path: Path,
    qc_form_type: str,
    gen_date: Optional[datetime] = None,
) -> Path:
    """Loads a QC template, populates all sheets for the given form type, and saves as .xlsx."""
    if not template_path.exists():
        raise FileNotFoundError(f"QC template not found: {template_path}")

    if gen_date is None:
        gen_date = datetime.now()

    formatted_date = format_header_date(gen_date)
    study_code = parsed_config.metadata.study_code
    full_version = parsed_config.metadata.full_version
    qc_form_name = (
        "Validation Check QC Form"
        if qc_form_type == "validation"
        else "Design and Specification QC Form"
    )

    wb = openpyxl.load_workbook(str(template_path))
    try:
        valid_qc_sheet_names = {
            cfg.qc_sheet_name for cfg in SHEET_CONFIGS if cfg.qc_form_type == qc_form_type
        }
        for s in list(wb.sheetnames):
            if s not in valid_qc_sheet_names:
                wb.remove(wb[s])

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            cfg = SHEET_CONFIGS_BY_QC_NAME.get(sheet_name)
            update_sheet_header_footer(
                ws, study_code, full_version, formatted_date, qc_form_name=qc_form_name
            )

            if cfg and cfg.qc_form_type == qc_form_type:
                parsed_sheet = parsed_config.sheets.get(sheet_name)
                if parsed_sheet:
                    populate_worksheet(ws, parsed_sheet, cfg, full_version)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output_path))
        inject_header_logo_and_formatting(output_path, template_path)
        return output_path
    finally:
        wb.close()


def generate_initial_qc_forms(
    parsed_config: ParsedViedocConfig,
    templates_dir: Path,
    output_dir: Path,
    gen_date: Optional[datetime] = None,
) -> Tuple[Path, Path]:
    """Generates both the Validation Check QC Form and Design and Specification QC Form."""
    if gen_date is None:
        gen_date = datetime.now()

    study_code = parsed_config.metadata.study_code
    full_version = parsed_config.metadata.full_version

    val_template = find_template_path(templates_dir, "validation")
    design_template = find_template_path(templates_dir, "design")

    val_output_name = generate_output_filename(
        study_code=study_code,
        qc_form_name="Validation Check QC Form",
        design_version=full_version,
        gen_date=gen_date,
    )
    design_output_name = generate_output_filename(
        study_code=study_code,
        qc_form_name="Design and Specification QC Form",
        design_version=full_version,
        gen_date=gen_date,
    )

    val_output_path = output_dir / val_output_name
    design_output_path = output_dir / design_output_name

    populate_qc_form(
        parsed_config=parsed_config,
        template_path=val_template,
        output_path=val_output_path,
        qc_form_type="validation",
        gen_date=gen_date,
    )

    populate_qc_form(
        parsed_config=parsed_config,
        template_path=design_template,
        output_path=design_output_path,
        qc_form_type="design",
        gen_date=gen_date,
    )

    return val_output_path, design_output_path