"""Delta & Amendment Engine for Viedoc EDC QC Automation.

Compares an existing (historical or Default Study baseline) QC form against a new
Viedoc Configuration Report, evaluates Critical Attributes, populates QC Method,
preserves historical QC signatures for unchanged elements, highlights changed tabs/rows,
and formats 'requires re-QC: [...]' comments per SOI-EDC-QC-001 v2.0.
"""

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple
import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.critical_attributes import evaluate_row_critical_attributes
from src.generator import (
    DATA_ALIGNMENT,
    DATA_BORDER,
    DATA_FONT,
    QC_TRAILING_COLUMNS,
    adjust_column_widths,
    ensure_six_qc_headers,
    set_cell_value_safe,
    update_sheet_header_footer,
)
from src.models import (
    ParsedSheet,
    SHEET_CONFIGS,
    SHEET_CONFIGS_BY_QC_NAME,
    SheetConfig,
)
from src.parser import (
    format_header_date,
    is_help_or_info_row,
    parse_viedoc_report,
)

COLUMNS_EXCLUDED_FROM_RE_QC: Dict[str, Set[str]] = {
    "Viedoc - Items and Groups": {
        "alert trackers",
    },
}

YELLOW_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
YELLOW_TAB_COLOR = "FFF2CC"
REMOVED_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
REMOVED_FONT = Font(name="Arial", size=10, strike=True, color="7F7F7F")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="000000")
HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

DIFF_OLD_COLOR = "FFC00000"
DIFF_OLD_FONT = InlineFont(rFont="Arial", sz=10, strike=True, color=DIFF_OLD_COLOR)
DIFF_NEW_FONT = InlineFont(rFont="Arial", sz=10, strike=None, color="FF000000")

# Recognize both current (6-col) and legacy (5-col) QC trailing headers when reading old QC files
ALL_KNOWN_QC_HEADERS = set(QC_TRAILING_COLUMNS) | {"Subject ID"}

PRIMARY_KEY_COLUMNS: Dict[str, List[str]] = {
    "Viedoc - Data Checks": ["OID"],
    "Viedoc - Alerts": ["Internal description of alert", "Action type"],
    "Viedoc - Items and Groups": ["Form ID", "Item group ID", "Item ID"],
    "Viedoc - Events": ["Study event ID"],
    "Viedoc - Activities": ["Study event ID", "Activity ID"],
    "Viedoc - Forms": ["Event ID", "Activity ID", "Form ID"],
    "Viedoc - Functions": ["Event ID", "Activity ID", "Form ID", "Item group ID", "Item ID", "Type"],
    "Viedoc - Code lists": ["OID", "Code value"],
}


@dataclass
class FieldDiff:
    """Represents a change in a single field between versions."""
    field_name: str
    old_value: Any
    new_value: Any


@dataclass
class RecordQCData:
    """Historical QC documentation sign-offs preserved from old QC workbook."""
    critical_attributes: Optional[str] = None
    qc_method: Optional[str] = None
    validation_date: Optional[str] = None
    validator: Optional[str] = None
    validation_status: Optional[str] = None
    comment: Optional[str] = None


@dataclass
class RowDelta:
    """Detailed diff result for a single row."""
    change_type: str  # 'NEW', 'MODIFIED', 'UNCHANGED', 'REMOVED'
    record_key: Tuple[str, ...]
    key_description: str
    sheet_name: str
    diffs: List[FieldDiff] = field(default_factory=list)
    new_row_vals: List[Any] = field(default_factory=list)
    old_row_vals: List[Any] = field(default_factory=list)
    qc_data: RecordQCData = field(default_factory=RecordQCData)


@dataclass
class OldQCRecord:
    """Data and signatures extracted from an existing QC form."""
    key: Tuple[str, ...]
    spec_values: Dict[str, Any]
    full_row: List[Any]
    qc_data: RecordQCData


@dataclass
class OldQCSheetData:
    """Sheet data extracted from an existing QC form."""
    sheet_name: str
    headers: List[str]
    records: Dict[Tuple[str, ...], OldQCRecord] = field(default_factory=dict)


def normalize_value(val: Any) -> str:
    """Normalizes cell values for reliable comparison."""
    if val is None:
        return ""
    return str(val).replace("\r\n", "\n").replace("\r", "\n").strip()


def parse_old_qc_sheet(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    sheet_name: str,
    key_col_names: List[str],
) -> OldQCSheetData:
    """Extracts specification data, keys, and QC signatures from an existing QC sheet.
    Supports both 6-column ('Critical Attributes', 'QC Method', ...) and legacy 5-column ('Subject ID', ...) files.
    """
    if ws.max_row < 1:
        return OldQCSheetData(sheet_name=sheet_name, headers=[])

    headers = [
        str(ws.cell(1, c).value).strip() if ws.cell(1, c).value is not None else ""
        for c in range(1, ws.max_column + 1)
    ]
    while headers and not headers[-1]:
        headers.pop()

    header_to_idx = {h: i for i, h in enumerate(headers)}

    crit_idx = header_to_idx.get("Critical Attributes")
    method_idx = header_to_idx.get("QC Method", header_to_idx.get("Subject ID"))
    date_idx = header_to_idx.get("Validation Date")
    val_idx = header_to_idx.get("Validator")
    status_idx = header_to_idx.get("Validation Status")
    comment_idx = header_to_idx.get("Comment")

    key_indices = [header_to_idx.get(k) for k in key_col_names]

    records: Dict[Tuple[str, ...], OldQCRecord] = {}

    def _safe_get(row_v: List[Any], idx: Optional[int]) -> Optional[Any]:
        return row_v[idx] if idx is not None and idx < len(row_v) else None

    for r in range(2, ws.max_row + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, len(headers) + 1)]
        if not any(v is not None and str(v).strip() != "" for v in row_vals):
            continue
        if is_help_or_info_row(row_vals):
            continue

        key_parts = [
            normalize_value(row_vals[k_idx]) if k_idx is not None and k_idx < len(row_vals) else ""
            for k_idx in key_indices
        ]
        record_key = tuple(key_parts)

        spec_dict = {}
        for h, idx in header_to_idx.items():
            if h not in ALL_KNOWN_QC_HEADERS and h != "Design Version":
                spec_dict[h] = row_vals[idx] if idx < len(row_vals) else None

        qc_data = RecordQCData(
            critical_attributes=_safe_get(row_vals, crit_idx),
            qc_method=_safe_get(row_vals, method_idx),
            validation_date=_safe_get(row_vals, date_idx),
            validator=_safe_get(row_vals, val_idx),
            validation_status=_safe_get(row_vals, status_idx),
            comment=_safe_get(row_vals, comment_idx),
        )

        records[record_key] = OldQCRecord(
            key=record_key,
            spec_values=spec_dict,
            full_row=row_vals,
            qc_data=qc_data,
        )

    return OldQCSheetData(
        sheet_name=sheet_name,
        headers=headers,
        records=records,
    )


def compare_sheet_delta(
    old_sheet: OldQCSheetData,
    new_sheet: ParsedSheet,
    key_col_names: List[str],
) -> List[RowDelta]:
    """Compares old QC sheet against new parsed sheet and categorizes each row."""
    deltas: List[RowDelta] = []
    new_headers = new_sheet.lowest_headers
    new_header_to_idx = {h: i for i, h in enumerate(new_headers)}

    new_key_indices = [new_header_to_idx[k] for k in key_col_names if k in new_header_to_idx]
    seen_new_keys: Set[Tuple[str, ...]] = set()

    for row_vals in new_sheet.rows:
        key_parts = [
            normalize_value(row_vals[idx]) if idx < len(row_vals) else ""
            for idx in new_key_indices
        ]
        key = tuple(key_parts)
        seen_new_keys.add(key)

        if key not in old_sheet.records:
            deltas.append(
                RowDelta(
                    change_type="NEW",
                    record_key=key,
                    key_description=" / ".join(key),
                    sheet_name=new_sheet.qc_sheet_name,
                    new_row_vals=row_vals,
                )
            )
        else:
            old_rec = old_sheet.records[key]
            diffs: List[FieldDiff] = []
            excluded_cols = {
                c.strip().lower()
                for c in COLUMNS_EXCLUDED_FROM_RE_QC.get(new_sheet.qc_sheet_name, set())
            }

            for h in new_headers:
                if h.strip().lower() in excluded_cols:
                    continue

                if h in old_rec.spec_values:
                    old_val = old_rec.spec_values[h]
                    new_idx = new_header_to_idx[h]
                    new_val = row_vals[new_idx] if new_idx < len(row_vals) else None

                    if normalize_value(old_val) != normalize_value(new_val):
                        diffs.append(
                            FieldDiff(field_name=h, old_value=old_val, new_value=new_val)
                        )

            change_type = "MODIFIED" if diffs else "UNCHANGED"
            deltas.append(
                RowDelta(
                    change_type=change_type,
                    record_key=key,
                    key_description=" / ".join(key),
                    sheet_name=new_sheet.qc_sheet_name,
                    diffs=diffs,
                    new_row_vals=row_vals,
                    qc_data=old_rec.qc_data,
                )
            )

    for old_key, old_rec in old_sheet.records.items():
        if old_key not in seen_new_keys:
            deltas.append(
                RowDelta(
                    change_type="REMOVED",
                    record_key=old_key,
                    key_description=" / ".join(old_key),
                    sheet_name=new_sheet.qc_sheet_name,
                    old_row_vals=old_rec.full_row,
                    qc_data=old_rec.qc_data,
                )
            )

    return deltas


def write_amendment_sheet(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    deltas: List[RowDelta],
    cfg: SheetConfig,
    full_version: str,
    new_headers: List[str],
) -> None:
    """Populates an amendment QC sheet with 6 QC columns, Critical Attributes,
    auto-populated 'Code review', yellow row/tab highlights, and 'requires re-QC: [...]' comments.
    """
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    total_cols = ensure_six_qc_headers(ws, len(new_headers), cfg.has_design_version_col)

    # Highlight worksheet tab yellow if any NEW or MODIFIED rows exist (CR-06)
    has_active_changes = any(d.change_type in ("NEW", "MODIFIED") for d in deltas)
    ws.sheet_properties.tabColor = YELLOW_TAB_COLOR if has_active_changes else None

    new_header_to_offset = {}
    col_start = 2 if cfg.has_design_version_col else 1
    for i, h in enumerate(new_headers):
        new_header_to_offset[h] = col_start + i

    current_row = 2

    for delta in deltas:
        row_fill = None
        row_font = DATA_FONT

        if delta.change_type in ("NEW", "MODIFIED"):
            row_fill = YELLOW_FILL
        elif delta.change_type == "REMOVED":
            row_fill = REMOVED_FILL
            row_font = REMOVED_FONT

        # 1. Col 1 (Design Version if applicable)
        col_offset = 1
        if cfg.has_design_version_col:
            cell = ws.cell(current_row, col_offset)
            set_cell_value_safe(cell, full_version)
            cell.font = row_font
            cell.border = DATA_BORDER
            cell.alignment = DATA_ALIGNMENT
            if row_fill:
                cell.fill = row_fill
            col_offset += 1

        # 2. Specification Values & Critical Attributes calculation
        crit_attrs: Optional[str] = None
        default_qc_method: Optional[str] = None

        if delta.change_type != "REMOVED":
            crit_attrs, default_qc_method = evaluate_row_critical_attributes(
                cfg.qc_sheet_name, new_headers, delta.new_row_vals
            )

            diff_by_offset = {}
            if delta.change_type == "MODIFIED":
                for diff in delta.diffs:
                    if diff.field_name in new_header_to_offset:
                        diff_by_offset[new_header_to_offset[diff.field_name]] = diff

            for val in delta.new_row_vals:
                cell = ws.cell(current_row, col_offset)

                if col_offset in diff_by_offset:
                    diff = diff_by_offset[col_offset]
                    old_str = str(diff.old_value) if diff.old_value is not None else ""
                    new_str = str(diff.new_value) if diff.new_value is not None else ""

                    clean_old = old_str.rstrip("\r\n")
                    clean_new = new_str.rstrip("\r\n")

                    if clean_old and clean_new:
                        rt = CellRichText()
                        rt.append(TextBlock(DIFF_OLD_FONT, clean_old + "\n"))
                        rt.append(TextBlock(DIFF_NEW_FONT, clean_new))
                        cell.value = rt
                    elif clean_old and not clean_new:
                        rt = CellRichText()
                        rt.append(TextBlock(DIFF_OLD_FONT, clean_old))
                        cell.value = rt
                    else:
                        set_cell_value_safe(cell, new_str)

                    old_display = str(diff.old_value) if diff.old_value is not None else "(empty)"
                    comment_msg = f"[v{full_version} Modified]\nPrevious: {old_display}"
                    cell.comment = Comment(comment_msg, "QC Automation")
                else:
                    set_cell_value_safe(cell, val)

                if not isinstance(cell.value, CellRichText):
                    cell.font = row_font
                cell.border = DATA_BORDER
                cell.alignment = DATA_ALIGNMENT
                if row_fill:
                    cell.fill = row_fill
                col_offset += 1
        else:
            old_vals = delta.old_row_vals
            start_idx = 1 if cfg.has_design_version_col and len(old_vals) > len(new_headers) else 0
            extracted_old_spec = []
            for i in range(len(new_headers)):
                val_idx = start_idx + i
                val = old_vals[val_idx] if val_idx < len(old_vals) else ""
                extracted_old_spec.append(val)
                cell = ws.cell(current_row, col_offset)
                set_cell_value_safe(cell, val)
                cell.font = row_font
                cell.border = DATA_BORDER
                cell.alignment = DATA_ALIGNMENT
                cell.fill = REMOVED_FILL
                col_offset += 1

            crit_attrs, _ = evaluate_row_critical_attributes(
                cfg.qc_sheet_name, new_headers, extracted_old_spec
            )

        # 3. 6 Trailing QC Columns:
        # [Critical Attributes, QC Method, Validation Date, Validator, Validation Status, Comment]
        qc = delta.qc_data
        if delta.change_type == "UNCHANGED":
            preserved_method = qc.qc_method if qc.qc_method else default_qc_method
            set_cell_value_safe(ws.cell(current_row, col_offset), crit_attrs)
            set_cell_value_safe(ws.cell(current_row, col_offset + 1), preserved_method)
            set_cell_value_safe(ws.cell(current_row, col_offset + 2), qc.validation_date)
            set_cell_value_safe(ws.cell(current_row, col_offset + 3), qc.validator)
            set_cell_value_safe(ws.cell(current_row, col_offset + 4), qc.validation_status)
            set_cell_value_safe(ws.cell(current_row, col_offset + 5), qc.comment)

        elif delta.change_type == "MODIFIED":
            # CR-02, CR-03, CR-05: Set Critical Attributes, default QC Method, and 'requires re-QC: [...]'
            set_cell_value_safe(ws.cell(current_row, col_offset), crit_attrs)
            set_cell_value_safe(ws.cell(current_row, col_offset + 1), default_qc_method)
            set_cell_value_safe(ws.cell(current_row, col_offset + 2), None)
            set_cell_value_safe(ws.cell(current_row, col_offset + 3), None)

            status_cell = ws.cell(current_row, col_offset + 4)
            set_cell_value_safe(status_cell, None)
            if qc.validator or qc.validation_status:
                status_cell.comment = Comment(
                    f"Prior QC: {qc.validation_status or 'N/A'} by {qc.validator or 'Unknown'} on {qc.validation_date or 'N/A'}",
                    "QC Automation",
                )

            changed_fields_str = ", ".join(d.field_name for d in delta.diffs)
            re_qc_comment = f"requires re-QC: {changed_fields_str}"
            comment_cell = ws.cell(current_row, col_offset + 5)
            set_cell_value_safe(comment_cell, re_qc_comment)
            if qc.comment:
                comment_cell.comment = Comment(f"Prior Comment: {qc.comment}", "QC Automation")

        elif delta.change_type == "NEW":
            set_cell_value_safe(ws.cell(current_row, col_offset), crit_attrs)
            set_cell_value_safe(ws.cell(current_row, col_offset + 1), default_qc_method)
            set_cell_value_safe(ws.cell(current_row, col_offset + 2), None)
            set_cell_value_safe(ws.cell(current_row, col_offset + 3), None)
            set_cell_value_safe(ws.cell(current_row, col_offset + 4), None)
            set_cell_value_safe(ws.cell(current_row, col_offset + 5), f"[Added in v{full_version}]")

        elif delta.change_type == "REMOVED":
            set_cell_value_safe(ws.cell(current_row, col_offset), qc.critical_attributes or crit_attrs)
            set_cell_value_safe(ws.cell(current_row, col_offset + 1), qc.qc_method)
            set_cell_value_safe(ws.cell(current_row, col_offset + 2), qc.validation_date)
            set_cell_value_safe(ws.cell(current_row, col_offset + 3), qc.validator)
            set_cell_value_safe(ws.cell(current_row, col_offset + 4), f"REMOVED in v{full_version}")
            set_cell_value_safe(ws.cell(current_row, col_offset + 5), qc.comment)

        for c_idx in range(col_offset, col_offset + 6):
            cell = ws.cell(current_row, c_idx)
            cell.font = row_font
            cell.border = DATA_BORDER
            cell.alignment = DATA_ALIGNMENT
            if row_fill:
                cell.fill = row_fill

        current_row += 1

    last_row = max(current_row - 1, 1)
    last_col_letter = get_column_letter(total_cols)
    ws.auto_filter.ref = f"A1:{last_col_letter}{last_row}"
    adjust_column_widths(ws, total_cols, last_row)


def create_change_log_sheet(
    wb: openpyxl.Workbook,
    all_deltas: Dict[str, List[RowDelta]],
    prev_version: str,
    new_version: str,
) -> None:
    """Inserts a dedicated 'Amendment Change Log' sheet detailing all changes."""
    log_sheet_name = "Amendment Change Log"
    if log_sheet_name in wb.sheetnames:
        ws = wb[log_sheet_name]
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet(title=log_sheet_name, index=0)

    headers = [
        "Sheet Name",
        "Change Type",
        "Record Key",
        "Changed Field",
        "Previous Value",
        "Current Value",
        "Previous Validator",
        "Previous QC Status",
    ]

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(1, col_idx, h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = DATA_BORDER
        cell.alignment = Alignment(vertical="center", horizontal="center")

    current_row = 2
    for sheet_name, deltas in all_deltas.items():
        for d in deltas:
            if d.change_type == "UNCHANGED":
                continue

            if d.change_type == "MODIFIED":
                for diff in d.diffs:
                    set_cell_value_safe(ws.cell(current_row, 1), sheet_name)
                    set_cell_value_safe(ws.cell(current_row, 2), d.change_type)
                    set_cell_value_safe(ws.cell(current_row, 3), d.key_description)
                    set_cell_value_safe(ws.cell(current_row, 4), diff.field_name)
                    set_cell_value_safe(ws.cell(current_row, 5), str(diff.old_value) if diff.old_value is not None else "")
                    set_cell_value_safe(ws.cell(current_row, 6), str(diff.new_value) if diff.new_value is not None else "")
                    set_cell_value_safe(ws.cell(current_row, 7), d.qc_data.validator or "")
                    set_cell_value_safe(ws.cell(current_row, 8), d.qc_data.validation_status or "")

                    for c in range(1, 9):
                        cell = ws.cell(current_row, c)
                        cell.font = DATA_FONT
                        cell.border = DATA_BORDER
                        cell.alignment = DATA_ALIGNMENT
                        cell.fill = YELLOW_FILL
                    current_row += 1
            else:
                set_cell_value_safe(ws.cell(current_row, 1), sheet_name)
                set_cell_value_safe(ws.cell(current_row, 2), d.change_type)
                set_cell_value_safe(ws.cell(current_row, 3), d.key_description)
                set_cell_value_safe(ws.cell(current_row, 4), "(Entire Record)")
                set_cell_value_safe(ws.cell(current_row, 5), "N/A" if d.change_type == "NEW" else f"Existed in v{prev_version}")
                set_cell_value_safe(ws.cell(current_row, 6), f"Added in v{new_version}" if d.change_type == "NEW" else "Removed")
                set_cell_value_safe(ws.cell(current_row, 7), d.qc_data.validator or "")
                set_cell_value_safe(ws.cell(current_row, 8), d.qc_data.validation_status or "")

                row_fill = YELLOW_FILL if d.change_type == "NEW" else REMOVED_FILL
                row_font = DATA_FONT if d.change_type == "NEW" else REMOVED_FONT
                for c in range(1, 9):
                    cell = ws.cell(current_row, c)
                    cell.font = row_font
                    cell.border = DATA_BORDER
                    cell.alignment = DATA_ALIGNMENT
                    cell.fill = row_fill
                current_row += 1

    last_row = max(current_row - 1, 1)
    ws.auto_filter.ref = f"A1:H{last_row}"
    adjust_column_widths(ws, 8, last_row)


def generate_amendment_qc_form(
    old_qc_path: Path,
    new_config_report_path: Path,
    template_path: Path,
    output_path: Path,
    qc_form_type: str,
    gen_date: Optional[datetime] = None,
) -> Tuple[Path, Dict[str, List[RowDelta]]]:
    """Compares old QC workbook against new Viedoc Config Report and generates updated QC workbook."""
    if gen_date is None:
        gen_date = datetime.now()

    new_config = parse_viedoc_report(new_config_report_path)
    new_version = new_config.metadata.full_version
    study_code = new_config.metadata.study_code
    formatted_date = format_header_date(gen_date)

    wb_old = openpyxl.load_workbook(str(old_qc_path), data_only=True)
    old_sheets: Dict[str, OldQCSheetData] = {}
    try:
        for s_name in wb_old.sheetnames:
            if s_name in PRIMARY_KEY_COLUMNS:
                keys = PRIMARY_KEY_COLUMNS[s_name]
                old_sheets[s_name] = parse_old_qc_sheet(wb_old[s_name], s_name, keys)
    finally:
        wb_old.close()

    prev_version = "Old"
    for s in old_sheets.values():
        if s.records:
            first_rec = next(iter(s.records.values()))
            if first_rec.full_row:
                val = str(first_rec.full_row[0]).strip()
                if "." in val:
                    prev_version = val
                    break

    all_deltas: Dict[str, List[RowDelta]] = {}
    for sheet_name, cfg in SHEET_CONFIGS_BY_QC_NAME.items():
        if cfg.qc_form_type == qc_form_type and sheet_name in new_config.sheets:
            new_sheet = new_config.sheets[sheet_name]
            old_sheet = old_sheets.get(sheet_name, OldQCSheetData(sheet_name=sheet_name, headers=[]))
            keys = PRIMARY_KEY_COLUMNS[sheet_name]
            deltas = compare_sheet_delta(old_sheet, new_sheet, keys)
            all_deltas[sheet_name] = deltas

    wb_out = openpyxl.load_workbook(str(template_path))
    try:
        valid_qc_sheet_names = {
            cfg.qc_sheet_name for cfg in SHEET_CONFIGS if cfg.qc_form_type == qc_form_type
        }
        for s in list(wb_out.sheetnames):
            if s not in valid_qc_sheet_names and s != "Amendment Change Log":
                wb_out.remove(wb_out[s])

        for sheet_name in wb_out.sheetnames:
            ws = wb_out[sheet_name]
            cfg = SHEET_CONFIGS_BY_QC_NAME.get(sheet_name)
            update_sheet_header_footer(ws, study_code, new_version, formatted_date)

            if cfg and cfg.qc_form_type == qc_form_type and sheet_name in all_deltas:
                deltas = all_deltas[sheet_name]
                new_headers = new_config.sheets[sheet_name].lowest_headers
                write_amendment_sheet(ws, deltas, cfg, new_version, new_headers)

        create_change_log_sheet(wb_out, all_deltas, prev_version, new_version)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb_out.save(str(output_path))
        return output_path, all_deltas
    finally:
        wb_out.close()