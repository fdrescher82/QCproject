"""Unit and integration tests for Phase 3: Delta & Amendment Engine."""

from datetime import datetime
from pathlib import Path
import openpyxl
import pytest

from src.delta_engine import generate_amendment_qc_form
from src.generator import find_template_path, generate_initial_qc_forms
from src.parser import parse_viedoc_report

TEMPLATES_DIR = Path(r"d:\QCproject\templates")
REPORT_20_PATH = Path(r"d:\QCproject\SurfSubQ ML45551_Configuration Report_20.0_20260302.xlsx")
REPORT_23_PATH = Path(r"d:\QCproject\SurfSubQ ML45551_Configuration Report_23.0_20260629.xlsx")
TEST_OUTPUT_DIR = Path(r"d:\QCproject\output_test")


def test_delta_engine_validation_form_amendment():
    """Simulates realistic study amendment workflow on Validation Check QC Form."""
    assert REPORT_20_PATH.exists()
    assert REPORT_23_PATH.exists()

    gen_date_20 = datetime(2026, 3, 5)
    parsed_config_20 = parse_viedoc_report(REPORT_20_PATH)

    val_v20_path, _ = generate_initial_qc_forms(
        parsed_config=parsed_config_20,
        templates_dir=TEMPLATES_DIR,
        output_dir=TEST_OUTPUT_DIR,
        gen_date=gen_date_20,
    )

    # Simulate tester filling out QC sign-offs on v20.0
    wb_val_v20 = openpyxl.load_workbook(str(val_v20_path))
    ws_dc = wb_val_v20["Viedoc - Data Checks"]

    # Sign off row 2 (RC_AENUM_23_0_1, which is UNCHANGED in v23.0)
    assert ws_dc.cell(2, 1).value == "RC_AENUM_23_0_1"
    ws_dc.cell(2, 9, "SUBJ-001")          # Subject ID
    ws_dc.cell(2, 10, "05/MAR/2026")      # Validation Date
    ws_dc.cell(2, 11, "Florian Drescher") # Validator
    ws_dc.cell(2, 12, "Passed")           # Validation Status
    ws_dc.cell(2, 13, "Verified in EDC")  # Comment

    # Sign off RC_BMI_8_0_1 (which is MODIFIED in v23.0)
    bmi_row_idx = None
    for r in range(2, ws_dc.max_row + 1):
        if ws_dc.cell(r, 1).value == "RC_BMI_8_0_1":
            bmi_row_idx = r
            break
    assert bmi_row_idx is not None, "RC_BMI_8_0_1 not found in v20"
    ws_dc.cell(bmi_row_idx, 9, "SUBJ-002")
    ws_dc.cell(bmi_row_idx, 10, "06/MAR/2026")
    ws_dc.cell(bmi_row_idx, 11, "Florian Drescher")
    ws_dc.cell(bmi_row_idx, 12, "Passed")
    ws_dc.cell(bmi_row_idx, 13, "BMI check passed")

    wb_val_v20.save(str(val_v20_path))
    wb_val_v20.close()

    # Run Amendment Engine
    gen_date_23 = datetime(2026, 10, 5)
    amendment_out_path = TEST_OUTPUT_DIR / "SurfSubQ ML45551_20261005_Validation Check QC Form_23.0.xlsx"
    val_template = find_template_path(TEMPLATES_DIR, "validation")

    out_file, all_deltas = generate_amendment_qc_form(
        old_qc_path=val_v20_path,
        new_config_report_path=REPORT_23_PATH,
        template_path=val_template,
        output_path=amendment_out_path,
        qc_form_type="validation",
        gen_date=gen_date_23,
    )

    assert out_file.exists()

    # Verify contents of Amendment QC Workbook
    wb_amend = openpyxl.load_workbook(str(out_file), data_only=False, rich_text=True)
    try:
        assert "Amendment Change Log" in wb_amend.sheetnames
        ws_log = wb_amend["Amendment Change Log"]
        assert ws_log.max_row > 1

        ws_dc_amend = wb_amend["Viedoc - Data Checks"]

        # 1. Unchanged row: signatures PRESERVED, no yellow fill
        assert ws_dc_amend.cell(2, 1).value == "RC_AENUM_23_0_1"
        assert ws_dc_amend.cell(2, 9).value == "SUBJ-001"
        assert ws_dc_amend.cell(2, 10).value == "05/MAR/2026"
        assert ws_dc_amend.cell(2, 11).value == "Florian Drescher"
        assert ws_dc_amend.cell(2, 12).value == "Passed"
        assert ws_dc_amend.cell(2, 13).value == "Verified in EDC"
        assert ws_dc_amend.cell(2, 1).fill.fill_type is None or ws_dc_amend.cell(2, 1).fill.start_color.rgb == "00000000"

        # 2. Modified row: yellow fill, cell rich text with prior value and new value, comment, Requires Re-QC
        bmi_amend_idx = None
        for r in range(2, ws_dc_amend.max_row + 1):
            if ws_dc_amend.cell(r, 1).value == "RC_BMI_8_0_1":
                bmi_amend_idx = r
                break
        assert bmi_amend_idx is not None
        assert ws_dc_amend.cell(bmi_amend_idx, 1).fill.start_color.rgb == "00FFF2CC"

        # Cell contains rich text with red strikethrough prior value and normal new value
        from openpyxl.cell.rich_text import CellRichText
        diff_cell = ws_dc_amend.cell(bmi_amend_idx, 6)
        assert isinstance(diff_cell.value, CellRichText)
        cell_val_str = str(diff_cell.value)
        assert "BMI >= 18.5 && BMI <= 30" in cell_val_str
        assert "BMI >= 15 && BMI <= 50" in cell_val_str

        # Check inline font formatting on parts
        old_part = diff_cell.value[0]
        assert old_part.font.strike is True
        assert old_part.font.color.rgb == "FFC00000"

        comment = diff_cell.comment
        assert comment is not None
        assert "BMI >= 18.5 && BMI <= 30" in comment.text

        assert ws_dc_amend.cell(bmi_amend_idx, 12).value == "Requires Re-QC"
        status_comment = ws_dc_amend.cell(bmi_amend_idx, 12).comment
        assert status_comment is not None
        assert "Florian Drescher" in status_comment.text

        # 3. Removed check: strike-through and gray fill
        removed_rows = [
            r for r in range(2, ws_dc_amend.max_row + 1)
            if ws_dc_amend.cell(r, 12).value == "REMOVED in v23.0"
        ]
        assert len(removed_rows) >= 1
        removed_idx = removed_rows[0]
        assert ws_dc_amend.cell(removed_idx, 1).font.strike is True
        assert ws_dc_amend.cell(removed_idx, 1).fill.start_color.rgb == "00F2F2F2"

        # 4. Page Header updated to 23.0 and uses black text (no red codes)
        hf = ws_dc_amend.HeaderFooter.oddHeader
        assert "23.0" in hf.center.text
        assert "&KFF0000" not in hf.center.text
        assert "&K01+000" not in hf.left.text
    finally:
        wb_amend.close()


def test_delta_engine_design_form_amendment():
    """Simulates study amendment workflow on Design and Specification QC Form."""
    assert REPORT_20_PATH.exists()
    assert REPORT_23_PATH.exists()

    gen_date_20 = datetime(2026, 3, 5)
    parsed_config_20 = parse_viedoc_report(REPORT_20_PATH)

    _, design_v20_path = generate_initial_qc_forms(
        parsed_config=parsed_config_20,
        templates_dir=TEMPLATES_DIR,
        output_dir=TEST_OUTPUT_DIR,
        gen_date=gen_date_20,
    )

    # Sign off first row in Items and Groups
    wb_des_v20 = openpyxl.load_workbook(str(design_v20_path))
    ws_items = wb_des_v20["Viedoc - Items and Groups"]
    # Row 2: Col 1 is Design Version, Col 2 is Form ID (AE), Col 64 is Subject ID
    assert ws_items.cell(2, 2).value == "AE"
    ws_items.cell(2, 64, "SUBJ-001")
    ws_items.cell(2, 65, "05/MAR/2026")
    ws_items.cell(2, 66, "Florian Drescher")
    ws_items.cell(2, 67, "Passed")
    ws_items.cell(2, 68, "Initial form structure approved")

    wb_des_v20.save(str(design_v20_path))
    wb_des_v20.close()

    # Run Amendment Engine on Design QC Form
    gen_date_23 = datetime(2026, 10, 5)
    amendment_design_out = TEST_OUTPUT_DIR / "SurfSubQ ML45551_20261005_Design and Specification QC Form_23.0.xlsx"
    design_template = find_template_path(TEMPLATES_DIR, "design")

    out_file, all_deltas = generate_amendment_qc_form(
        old_qc_path=design_v20_path,
        new_config_report_path=REPORT_23_PATH,
        template_path=design_template,
        output_path=amendment_design_out,
        qc_form_type="design",
        gen_date=gen_date_23,
    )

    assert out_file.exists()

    wb_amend = openpyxl.load_workbook(str(out_file), data_only=False)
    try:
        # Check Change Log exists
        assert "Amendment Change Log" in wb_amend.sheetnames
        ws_log = wb_amend["Amendment Change Log"]
        assert ws_log.max_row > 10, "Expected multiple design changes logged"

        # Check Items and Groups
        ws_items_amend = wb_amend["Viedoc - Items and Groups"]
        # Row 2 (Unchanged item): Sign-off preserved!
        assert ws_items_amend.cell(2, 1).value == "23.0"
        assert ws_items_amend.cell(2, 64).value == "SUBJ-001"
        assert ws_items_amend.cell(2, 65).value == "05/MAR/2026"
        assert ws_items_amend.cell(2, 66).value == "Florian Drescher"
        assert ws_items_amend.cell(2, 67).value == "Passed"

        # Check Functions: contains modified and removed functions
        ws_funcs = wb_amend["Viedoc - Functions"]
        assert ws_funcs.max_row > 100

        # Check Code lists: 747 items (with 7 new yellow rows)
        ws_codes = wb_amend["Viedoc - Code lists"]
        assert ws_codes.max_row >= 748
    finally:
        wb_amend.close()


def test_alert_trackers_excluded_from_triggering_re_qc():
    """Verifies that changes to 'Alert trackers' column do NOT trigger re-QC (Issue #4)."""
    from src.delta_engine import compare_sheet_delta, OldQCRecord, OldQCSheetData, RecordQCData
    from src.models import ParsedSheet

    sheet_name = "Viedoc - Items and Groups"
    headers = ["Form ID", "Item group ID", "Item ID", "Field label", "Alert trackers"]
    keys = ["Form ID", "Item group ID", "Item ID"]

    # Historical record 1: Only Alert trackers will change
    rec1_spec = {
        "Form ID": "DM", "Item group ID": "DM_GRP", "Item ID": "AGE",
        "Field label": "Age in years", "Alert trackers": ""
    }
    rec1_qc = RecordQCData(validator="Florian Drescher", validation_status="Passed")
    old_rec1 = OldQCRecord(
        key=("DM", "DM_GRP", "AGE"),
        spec_values=rec1_spec,
        qc_data=rec1_qc,
        full_row=["DM", "DM_GRP", "AGE", "Age in years", ""]
    )

    # Historical record 2: Field label AND Alert trackers will change
    rec2_spec = {
        "Form ID": "VS", "Item group ID": "VS_GRP", "Item ID": "SYSBP",
        "Field label": "Systolic BP", "Alert trackers": ""
    }
    rec2_qc = RecordQCData(validator="Florian Drescher", validation_status="Passed")
    old_rec2 = OldQCRecord(
        key=("VS", "VS_GRP", "SYSBP"),
        spec_values=rec2_spec,
        qc_data=rec2_qc,
        full_row=["VS", "VS_GRP", "SYSBP", "Systolic BP", ""]
    )

    old_sheet = OldQCSheetData(
        sheet_name=sheet_name,
        headers=headers,
        records={
            ("DM", "DM_GRP", "AGE"): old_rec1,
            ("VS", "VS_GRP", "SYSBP"): old_rec2,
        }
    )

    # New sheet:
    # Row 1 has ONLY Alert trackers changed ("" -> "ALERT_AGE_CHECK")
    # Row 2 has Field label changed ("Systolic BP" -> "Systolic Blood Pressure") and Alert trackers changed
    new_rows = [
        ["DM", "DM_GRP", "AGE", "Age in years", "ALERT_AGE_CHECK"],
        ["VS", "VS_GRP", "SYSBP", "Systolic Blood Pressure", "ALERT_BP_HIGH"],
    ]
    new_sheet = ParsedSheet(
        qc_sheet_name=sheet_name,
        viedoc_sheet_name="Items and Groups",
        lowest_headers=headers,
        rows=new_rows
    )

    deltas = compare_sheet_delta(old_sheet, new_sheet, keys)

    assert len(deltas) == 2
    delta_age = next(d for d in deltas if d.record_key == ("DM", "DM_GRP", "AGE"))
    delta_bp = next(d for d in deltas if d.record_key == ("VS", "VS_GRP", "SYSBP"))

    # AGE item must be UNCHANGED (Alert trackers changes do NOT trigger re-QC)
    assert delta_age.change_type == "UNCHANGED", "Alert trackers change should not trigger re-QC"
    assert delta_age.diffs == []

    # BP item must be MODIFIED (because Field label changed)
    assert delta_bp.change_type == "MODIFIED"
    assert len(delta_bp.diffs) == 1
    assert delta_bp.diffs[0].field_name == "Field label"

